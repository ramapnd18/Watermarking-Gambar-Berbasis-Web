"""jalankan_matriks.py — Runner matriks pengujian UJI.md §3 + §6 → XLSX.

Menjalankan seluruh kombinasi uji utama secara batch (lebih cepat dan
reproduksibel dibanding klik manual di UI):
- §3.1 PSNR embed untuk 3 citra (logo pola 32x32, alpha 15).
- §3.2 JPEG 90/70/50 + ekstraksi NC/BER per citra (9 kombinasi).
- §3.3 noise/crop-25/resize/kontras + gabungan + ekstraksi (15 kombinasi).
- §6  trade-off: citra_02 × alpha 5/15/30 → PSNR + NC/BER pasca-JPEG50.

Masukan : tests/assets/citra_01_tekstur_tinggi.*,
          tests/assets/citra_02_area_datar.*,
          tests/assets/citra_03_potret_wajah.* (jpg/png).
Keluaran: hasil_pengujian.xlsx (sheet PSNR, Serangan_JPEG,
          Serangan_Manipulasi, Tradeoff_Alpha) + citra hasil di
          tests/assets/hasil/<citra>/ untuk reproduksi dosen.

Kunci uji fixed per citra ("kunci-uji-<nama>") agar hasil reproduksibel;
INI BUKAN kunci produksi (lihat security.py untuk kunci asli via secrets).

Cara pakai (dari root, venv aktif):
    venv\\Scripts\\python tools/jalankan_matriks.py
"""

import sys
import time
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font
from PIL import Image

AKAR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(AKAR))

from backend.attacker import (
    kompresi_jpeg, serangan_noise_gaussian, serangan_crop_tengah,
    serangan_resize, serangan_kontras, serangan_q50_noise_crop,
)
from backend.evaluator import hitung_psnr, hitung_nc, hitung_ber
from backend.watermark_engine import (
    sisip_watermark, ekstrak_watermark, logo_ke_bit,
)

CITRA_UJI = (
    "citra_01_tekstur_tinggi",
    "citra_02_area_datar",
    "citra_03_potret_wajah",
)
# Slot pengayaan (§5, opsional): bila berkas ini ADA, matriks mini dijalankan
# otomatis (embed + jpeg50/noise/crop + ekstrak) dan dibandingkan dengan
# rata-rata 3 citra utama di sheet Pengayaan_AI. Bila TIDAK ada, dilewati
# tanpa error — pengayaan bukan syarat kelulusan.
CITRA_PENGAYAAN = "pengayaan/citra_ai_generatif"
SERANGAN_PENGAYAAN = ("jpeg_50", "noise_gaussian", "crop_tengah_25")
ALPHA_MATRIKS = 15.0
ALPHA_TRADEOFF = (5.0, 15.0, 30.0)
LOGO_PATH = AKAR / "assets" / "logo_pola_32x32.png"
DIR_ASET = AKAR / "tests" / "assets"
DIR_HASIL = DIR_ASET / "hasil"
XLSX_KELUAR = AKAR / "hasil_pengujian.xlsx"


def cari_berkas_citra(nama: str) -> Path:
    """Mencari file citra uji dengan berbagai ekstensi umum."""
    for ekst in (".jpg", ".jpeg", ".png", ".JPG", ".PNG"):
        kandidat = DIR_ASET / (nama + ekst)
        if kandidat.exists():
            return kandidat
    raise FileNotFoundError(
        f"Citra uji '{nama}.*' tidak ditemukan di {DIR_ASET}. "
        "Taruh 3 foto uji (tekstur tinggi, area datar, potret wajah) di sana, "
        "lalu jalankan ulang skrip ini."
    )


def fmt_psnr(nilai) -> str:
    """Format PSNR untuk sel XLSX (None -> 'n/a (ukuran berubah)')."""
    if nilai is None:
        return "n/a (ukuran berubah)"
    if nilai == float("inf"):
        return "inf (identik)"
    return round(float(nilai), 2)


def psnr_aman(asli: Image.Image, uji: Image.Image):
    """PSNR bila seukuran, else None (crop/gabungan, anti inf-semu)."""
    if asli.size != uji.size:
        return None
    a = np.asarray(asli.convert("RGB"), dtype=np.float64)
    u = np.asarray(uji.convert("RGB"), dtype=np.float64)
    mse = float(np.mean((a - u) ** 2))
    if mse == 0.0:
        return float("inf")
    return float(10.0 * np.log10((255.0 ** 2) / mse))


def ekstrak_aman(cover, citra_uji, kunci, jumlah_bit, alpha):
    """Ekstraksi + NC/BER; gagal sinkron -> baris jujur 'tak-tersinkron'."""
    try:
        bit_hasil = ekstrak_watermark(cover, citra_uji, kunci, jumlah_bit, alpha)
    except ValueError as eronya:
        return {"nc": "tak-tersinkron", "ber": "tak-tersinkron",
                "status": str(eronya)[:120]}
    nc = hitung_nc(BIT_ACUAN, bit_hasil)
    ber = hitung_ber(BIT_ACUAN, bit_hasil)
    if nc >= 0.75:
        status = "Watermark Terdeteksi Kuat"
    elif nc >= 0.4:
        status = "Terdeteksi Sebagian"
    else:
        status = "Tidak Terdeteksi"
    return {"nc": round(nc, 4), "ber": round(ber * 100, 2), "status": status}


SERANGAN = (
    ("jpeg_90", lambda img: kompresi_jpeg(img, kualitas=90)),
    ("jpeg_70", lambda img: kompresi_jpeg(img, kualitas=70)),
    ("jpeg_50", lambda img: kompresi_jpeg(img, kualitas=50)),
    ("noise_gaussian", lambda img: serangan_noise_gaussian(img, seed=1)),
    ("crop_tengah_25", lambda img: serangan_crop_tengah(img, proporsi=0.25)),
    ("resize_05", lambda img: serangan_resize(img, skala=0.5)),
    ("kontras_12", lambda img: serangan_kontras(img, faktor=1.2, kecerahan=20.0)),
    ("gabungan_q50_noise_crop",
     lambda img: serangan_q50_noise_crop(img, seed=1)),
)


def tulis_sheet(wb, judul, kepala, baris):
    """Membuat satu sheet dengan header tebal dari daftar dict/baris."""
    ws = wb.create_sheet(judul)
    for kolom, nama in enumerate(kepala, start=1):
        sel = ws.cell(row=1, column=kolom, value=nama)
        sel.font = Font(bold=True)
    for i, b in enumerate(baris, start=2):
        for kolom, nama in enumerate(kepala, start=1):
            ws.cell(row=i, column=kolom, value=b.get(nama))
    for kolom in range(1, len(kepala) + 1):
        ws.column_dimensions[ws.cell(row=1, column=kolom).column_letter].width = 22
    return ws


def jalankan():
    """Orkestrasi seluruh matriks + tulis XLSX. Return 0 bila sukses."""
    print("[matriks] Memuat logo acuan:", LOGO_PATH.name)
    global BIT_ACUAN
    bit_logo, lebar_logo, tinggi_logo = logo_ke_bit(Image.open(LOGO_PATH))
    BIT_ACUAN = bit_logo
    print(f"[matriks] Logo {lebar_logo}x{tinggi_logo} = {len(bit_logo)} bit.")

    baris_psnr, baris_jpeg, baris_manip = [], [], []
    for nama in CITRA_UJI:
        berkas = cari_berkas_citra(nama)
        cover = Image.open(berkas)
        cover.load()
        kunci = f"kunci-uji-{nama}"
        dir_citra = DIR_HASIL / nama
        dir_citra.mkdir(parents=True, exist_ok=True)
        print(f"[matriks] {nama}: {cover.size[0]}x{cover.size[1]} {cover.mode}")

        t0 = time.time()
        stego = sisip_watermark(cover, bit_logo, kunci, alpha=ALPHA_MATRIKS)
        stego.save(dir_citra / f"stego_alpha{ALPHA_MATRIKS:g}.png")
        psnr = psnr_aman(cover, stego)
        baris_psnr.append({"Citra": nama, "Alpha": ALPHA_MATRIKS,
                           "PSNR (dB)": fmt_psnr(psnr),
                           "Waktu (s)": round(time.time() - t0, 1)})

        for nama_serang, fungsi in SERANGAN:
            t1 = time.time()
            hasil = fungsi(stego)
            hasil.save(dir_citra / f"{nama_serang}.png")
            met = ekstrak_aman(cover, hasil, kunci, len(bit_logo), ALPHA_MATRIKS)
            baris = {"Citra": nama, "Serangan": nama_serang,
                     "PSNR (dB)": fmt_psnr(psnr_aman(stego, hasil)),
                     "NC": met["nc"], "BER (%)": met["ber"],
                     "Status": met["status"],
                     "Waktu (s)": round(time.time() - t1, 1)}
            (baris_jpeg if nama_serang.startswith("jpeg_") else baris_manip).append(baris)
            print(f"  {nama_serang}: PSNR={baris['PSNR (dB)']} NC={met['nc']} BER={met['ber']}")

    # §6 trade-off: citra_02 × 3 alpha -> PSNR + NC/BER pasca-JPEG50.
    print("[matriks] Trade-off alpha (citra_02, serangan JPEG 50)...")
    baris_trade = []
    cover2 = Image.open(cari_berkas_citra("citra_02_area_datar"))
    cover2.load()
    kunci2 = "kunci-uji-citra_02_area_datar"
    for alpha in ALPHA_TRADEOFF:
        stego = sisip_watermark(cover2, bit_logo, kunci2, alpha=alpha)
        stego.save(DIR_HASIL / "citra_02_area_datar" / f"stego_alpha{alpha:g}.png")
        uji = kompresi_jpeg(stego, kualitas=50)
        met = ekstrak_aman(cover2, uji, kunci2, len(bit_logo), alpha)
        baris_trade.append({"Alpha": alpha,
                            "PSNR embed (dB)": fmt_psnr(psnr_aman(cover2, stego)),
                            "NC pasca-JPEG50": met["nc"],
                            "BER pasca-JPEG50 (%)": met["ber"]})
        print(f"  alpha={alpha}: PSNR={baris_trade[-1]['PSNR embed (dB)']} NC={met['nc']}")

    wb = Workbook()
    tulis_sheet(wb, "PSNR", ("Citra", "Alpha", "PSNR (dB)", "Waktu (s)"), baris_psnr)
    tulis_sheet(wb, "Serangan_JPEG",
                ("Citra", "Serangan", "PSNR (dB)", "NC", "BER (%)", "Status", "Waktu (s)"),
                baris_jpeg)
    tulis_sheet(wb, "Serangan_Manipulasi",
                ("Citra", "Serangan", "PSNR (dB)", "NC", "BER (%)", "Status", "Waktu (s)"),
                baris_manip)
    tulis_sheet(wb, "Tradeoff_Alpha",
                ("Alpha", "PSNR embed (dB)", "NC pasca-JPEG50", "BER pasca-JPEG50 (%)"),
                baris_trade)

    # Slot pengayaan §5: hanya bila gambar AI disediakan pengguna.
    try:
        berkas_ai = cari_berkas_citra(CITRA_PENGAYAAN)
    except FileNotFoundError:
        berkas_ai = None
    if berkas_ai is not None:
        print("[matriks] Pengayaan: citra AI ditemukan, menjalankan matriks mini...")
        baris_ai = _matriks_pengayaan(berkas_ai, baris_jpeg, baris_manip)
        tulis_sheet(wb, "Pengayaan_AI",
                    ("Serangan", "NC (AI)", "BER (AI) %", "NC rata-rata 3 citra",
                     "BER rata-rata 3 citra %", "Selisih NC"),
                    baris_ai)
    else:
        print("[matriks] Pengayaan dilewati (tidak ada tests/assets/pengayaan/citra_ai_generatif.*).")

    wb.remove(wb["Sheet"])
    wb.save(XLSX_KELUAR)
    print("[matriks] Selesai ->", XLSX_KELUAR)
    return 0


def _matriks_pengayaan(berkas_ai: Path, baris_jpeg: list, baris_manip: list) -> list:
    """Matriks mini §5 + perbandingan vs rata-rata 3 citra utama."""
    cover = Image.open(berkas_ai)
    cover.load()
    kunci = "kunci-uji-pengayaan-ai"
    dir_ai = DIR_HASIL / "pengayaan"
    dir_ai.mkdir(parents=True, exist_ok=True)
    stego = sisip_watermark(cover, BIT_ACUAN, kunci, alpha=ALPHA_MATRIKS)
    stego.save(dir_ai / "stego_ai.png")
    peta = dict(SERANGAN)
    baris = []
    for nama_serang in SERANGAN_PENGAYAAN:
        hasil = peta[nama_serang](stego)
        hasil.save(dir_ai / f"{nama_serang}_ai.png")
        met = ekstrak_aman(cover, hasil, kunci, len(BIT_ACUAN), ALPHA_MATRIKS)
        pembanding = (baris_jpeg if nama_serang.startswith("jpeg_") else baris_manip)
        nc_utama = [b["NC"] for b in pembanding
                    if b["Serangan"] == nama_serang and isinstance(b["NC"], float)]
        ber_utama = [b["BER (%)"] for b in pembanding
                     if b["Serangan"] == nama_serang and isinstance(b["BER (%)"], float)]
        nc_avg = round(sum(nc_utama) / len(nc_utama), 4)
        ber_avg = round(sum(ber_utama) / len(ber_utama), 2)
        baris.append({
            "Serangan": nama_serang,
            "NC (AI)": met["nc"], "BER (AI) %": met["ber"],
            "NC rata-rata 3 citra": nc_avg, "BER rata-rata 3 citra %": ber_avg,
            "Selisih NC": round(met["nc"] - nc_avg, 4)
            if isinstance(met["nc"], float) else "tak-tersinkron",
        })
        print(f"  AI {nama_serang}: NC={met['nc']} (utama {nc_avg})")
    return baris


if __name__ == "__main__":
    try:
        sys.exit(jalankan())
    except FileNotFoundError as eronya:
        print("[matriks] GAGAL:", eronya)
        sys.exit(2)
